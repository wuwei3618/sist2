#!/usr/bin/env python3
"""Regenerates the archive and legacy Office fixtures in this directory.

Needs: bsdtar (libarchive-tools), rar (rarlab.com), LibreOffice (soffice), python-pptx, openpyxl.
    python3 generate.py

The tests assert on the words written here, so keep them in sync with ArchiveTests/OLETests.
"""
import gzip
import io
import os
import shutil
import struct
import subprocess
import tarfile
import tempfile
import zipfile
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))


def out(name):
    return os.path.join(HERE, name)


def write_tree(root):
    files = {
        "docs/说明.txt": "压缩包里的中文说明 archive readme",
        "docs/notes.md": "tar member mentions walrus",
    }
    for path, text in files.items():
        full = os.path.join(root, path)
        os.makedirs(os.path.dirname(full), exist_ok=True)
        with open(full, "w", encoding="utf-8") as f:
            f.write(text)
    return files


def inner_zip():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("deep.txt", "nested zip inside tar mentions narwhal")
    return buffer.getvalue()


def stored_zip(entries):
    """A zip whose names are raw bytes without the UTF-8 flag, as old Windows tools write them."""
    data_out, central = b"", b""
    for name, data in entries:
        crc = zlib.crc32(data) & 0xFFFFFFFF
        offset = len(data_out)
        data_out += struct.pack("<IHHHHHIIIHH", 0x04034B50, 20, 0, 0, 0, 0x21, crc, len(data), len(data), len(name), 0)
        data_out += name + data
        central += struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, 20, 20, 0, 0, 0, 0x21, crc, len(data), len(data),
                               len(name), 0, 0, 0, 0, 0, offset) + name
    end = struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, len(entries), len(entries), len(central), len(data_out), 0)
    return data_out + central + end


def png(width, height, rgb):
    def chunk(kind, body):
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
    rows = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


def archives(tmp):
    tree = os.path.join(tmp, "tree")
    write_tree(tree)

    # tar.gz (pax headers carry the UTF-8 names), with a zip inside
    with tarfile.open(out("archive.tar.gz"), "w:gz", format=tarfile.PAX_FORMAT) as tar:
        tar.add(os.path.join(tree, "docs"), arcname="docs")
        data = inner_zip()
        info = tarfile.TarInfo("inner.zip")
        info.size = len(data)
        info.mtime = 1_700_000_000
        tar.addfile(info, io.BytesIO(data))

    subprocess.run(["bsdtar", "--format", "7zip", "-cf", out("archive.7z"), "-C", tree, "docs"], check=True)

    if os.path.exists(out("archive.rar")):
        os.remove(out("archive.rar"))
    subprocess.run(["rar", "a", "-ma5", "-r", "-idq", out("archive.rar"), "docs"], cwd=tree, check=True)

    with gzip.open(out("notes.txt.gz"), "wb") as f:
        f.write("single gzip file mentions pelican".encode())

    with open(out("gbk-names.zip"), "wb") as f:
        f.write(stored_zip([("中文名.txt".encode("gbk"), "gbk named member mentions albatross".encode())]))

    comic = os.path.join(tmp, "comic")
    os.makedirs(comic)
    with open(os.path.join(comic, "002.png"), "wb") as f:
        f.write(png(16, 24, (0, 0, 255)))
    with open(os.path.join(comic, "001.png"), "wb") as f:
        f.write(png(16, 24, (255, 0, 0)))
    if os.path.exists(out("comic.cbr")):
        os.remove(out("comic.cbr"))
    subprocess.run(["rar", "a", "-ma5", "-idq", out("comic.cbr"), "002.png", "001.png"], cwd=comic, check=True)


def zipped(path):
    """Large fixtures are kept deflated (sample.ppt holds a bitmap preview LibreOffice embeds)."""
    with zipfile.ZipFile(path + ".zip", "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        z.write(path, os.path.basename(path))
    os.remove(path)


def convert(source, extension, tmp):
    subprocess.run(["soffice", "--headless", "--convert-to", extension, "--outdir", tmp, source],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.join(tmp, os.path.splitext(os.path.basename(source))[0] + "." + extension)


def office(tmp):
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet["A1"] = "季度销售报告"
    sheet["A2"] = "华东地区"
    sheet["B2"] = 1234.5
    summary = book.create_sheet("Summary")
    summary["A1"] = "Quarterly revenue grew"
    book.properties.title = "测试工作簿"
    book.properties.creator = "张三"
    book.save(os.path.join(tmp, "sample.xlsx"))
    shutil.copy(convert(os.path.join(tmp, "sample.xlsx"), "xls", tmp), out("sample.xls"))

    # Enough unique strings for the shared string table to spill into CONTINUE records, with
    # Chinese (two-byte) and Latin (one-byte) strings alternating across the boundaries
    book = openpyxl.Workbook()
    sheet = book.active
    for i in range(1, 3001):
        sheet.cell(row=i, column=1, value=f"条目{i:04d} 中文内容" if i % 2 else f"entry {i:04d} latin text")
    sheet.cell(row=3001, column=1, value="最后一个条目 final entry")
    book.save(os.path.join(tmp, "many-strings.xlsx"))
    shutil.copy(convert(os.path.join(tmp, "many-strings.xlsx"), "xls", tmp), out("many-strings.xls"))
    zipped(out("many-strings.xls"))

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "年度总结"
    slide.placeholders[1].text = "Annual review"
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "第二页"
    slide.placeholders[1].text = "企鹅数量增长 penguin growth"
    slide.notes_slide.notes_text_frame.text = "演讲者备注 speaker notes"
    box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(1))
    box.text_frame.text = "文本框内容 textbox words"
    deck.core_properties.title = "测试演示文稿"
    deck.core_properties.author = "李四"
    deck.save(os.path.join(tmp, "sample.pptx"))
    shutil.copy(convert(os.path.join(tmp, "sample.pptx"), "ppt", tmp), out("sample.ppt"))
    zipped(out("sample.ppt"))


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        archives(tmp)
        office(tmp)
    print("fixtures written to", HERE)
